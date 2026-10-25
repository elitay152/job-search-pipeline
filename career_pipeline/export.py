import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from .models import Career, Packet
from .pipeline import audit
from .store import digest, now
from .grounding import accepted_assessments

from .text_safety import XML_INVALID
from .baseline import employment_baseline, evidence_roles
from types import SimpleNamespace


def require_clean_final_text(packet, career):
    if any(XML_INVALID.search(text) for text in preview(packet, career)):
        raise ValueError('Document text contains invalid control characters. Download a labeled draft and audit for inspection; correct the source or regenerate before final approval.')


def resume_title(evidence):
    return (evidence.resume_display_title
            if evidence.resume_display_title and evidence.resume_display_title_approved
            else evidence.official_title)


def resume_date(value):
    if not value:
        return 'Present'
    if isinstance(value, str):
        from datetime import date
        # The day is only a formatting aid, never stored or exported as a career fact.
        return date(int(value[:4]), int(value[5:7]), 1).strftime('%b %Y')
    return value.strftime('%b %Y')


def preview(packet, career):
    index = {e.id: e for e in career.evidence}
    baseline = employment_baseline(career)
    roles = {j['id']: SimpleNamespace(**j) for j in baseline['jobs']}
    resume = [f'# {career.display_name}', *career.contact]
    if packet.skills:
        categories = {}
        seen = set()
        for ident in packet.skills:
            e = index.get(ident)
            if not e or not e.eligible:
                continue
            for tech in e.technologies:
                key = (tech.name.casefold(), tech.context)
                if key in seen:
                    continue
                seen.add(key)
                label = tech.name if tech.context == 'professional' else f'{tech.name} ({tech.context})'
                categories.setdefault(tech.category, []).append(label)
        if categories:
            resume += ['## Technical Skills'] + [f'{name}: ' + ', '.join(items) for name, items in categories.items()]
    job_claims = {ident: [] for ident in roles}
    shared_claims = {}
    unplaced = []
    sections = {'project': [], 'education': [], 'publication': [], 'skill': []}
    for claim in packet.resume:
        e = index.get(claim.evidence_ids[0]) if claim.evidence_ids else None
        if e and e.kind != 'experience':
            sections[e.kind].append(claim)
            continue
        linked = [index[i] for i in claim.evidence_ids if i in index and index[i].kind == 'experience']
        scope = tuple(sorted({ident for item in linked for ident in evidence_roles(item, baseline)}))
        if len(scope) == 1 and scope[0] in job_claims:
            job_claims[scope[0]].append(claim)
        elif scope:
            shared_claims.setdefault(scope, []).append(claim)
        else:
            unplaced.append(claim)
    shared_labels = {}
    for scope, claims in shared_claims.items():
        jobs = [roles[i] for i in scope if i in roles]
        label = 'Shared across ' + '; '.join(f'{resume_title(j)} at {j.employer}' for j in jobs)
        for claim in claims:
            shared_labels[claim.id] = label
            for ident in scope:
                if ident in job_claims:
                    job_claims[ident].append(claim)
    if roles or unplaced:
        resume += ['## Professional Experience']
    # Headers come only from saved history, never the model's chosen bullets or their order.
    for ident, job in roles.items():
        resume += [f'### {resume_title(job)} — {job.employer}']
        if job.start:
            resume += [f'{resume_date(job.start)} – {resume_date(job.end)}']
        for claim in job_claims[ident]:
            scope_label = shared_labels.get(claim.id)
            resume += [f'- {scope_label}: {claim.text}' if scope_label else f'- {claim.text}']
    if unplaced:
        resume += ['### Experience requiring role clarification']
        resume += [f'- {claim.text}' for claim in unplaced]
    for kind, heading in [('project', 'Selected Projects'), ('education', 'Education'), ('publication', 'Publications')]:
        if sections[kind]:
            resume += ['## ' + heading]
            resume += [c.text for c in sections[kind]]
    letter = [f'# {career.display_name}', *career.contact, 'Dear Hiring Team,',
              *[c.text for c in packet.cover_letter], 'Thank you for your consideration.']
    return '\n\n'.join(resume), '\n\n'.join(letter) if packet.cover_letter else ''


def packet_audit(packet, career, manifest):
    accepted = set()
    if manifest.get('settings', {}).get('factual_mode') == 'model_reviewed':
        if manifest['settings'].get('reviewer') not in ('anthropic', 'openai'):
            raise ValueError('Mock reviewers cannot authorize paraphrases')
        accepted, _ = accepted_assessments(packet, career, manifest['versions'][-1].get('support_assessments', []))
    return audit(packet, career, accepted, include_cover_letter=manifest.get('settings', {}).get('include_cover_letter', True))


def require_review_complete(manifest):
    if manifest['status'] == 'failed' or not manifest['versions']:
        raise ValueError('No completed, reviewed packet')
    final = manifest['versions'][-1]
    if any(final.get('reviews', {}).get(role, {}).get('status') != 'passed'
           for role in ('evidence', 'recruiter', 'voice')):
        raise ValueError('All three reviews must complete before approval or final export')
    if any(f['severity'] == 'error' for f in final.get('findings', [])):
        raise ValueError('Unresolved factual/reviewer errors block approval and final export')


def approve(store, run_id, career, acknowledged=False):
    if not acknowledged:
        raise ValueError('Explicit user review acknowledgement required')
    m = store.get(run_id)
    if digest(career.model_dump(mode='json')) != m['evidence_hash']:
        raise ValueError('Evidence snapshot changed; rerun required')
    if not m['versions'] or m['status'] == 'failed':
        raise ValueError('No completed packet')
    require_review_complete(m)
    packet = Packet.model_validate(m['versions'][-1]['packet'])
    if packet_audit(packet, career, m):
        raise ValueError('Factual validation blocks approval')
    require_clean_final_text(packet, career)
    m.update(status='approved', approved=True, approved_at=now(), approved_packet_hash=digest(packet.model_dump()))
    store.save(m)
    return m

def export(store, run_id, career, directory='artifacts', draft=False):
    m = store.get(run_id)
    packet = Packet.model_validate(m['versions'][-1]['packet'])
    if digest(career.model_dump(mode='json')) != m['evidence_hash']:
        raise ValueError('Evidence snapshot mismatch')
    if any(f.code in ('missing_employment_history', 'missing_source_bullets') for f in audit(packet, career)):
        raise ValueError('Resume has a job without bullets. Create a new run using your source history before exporting.')
    if not draft:
        require_review_complete(m)
    if not draft and (packet_audit(packet, career, m) or not m.get('approved') or m.get('approved_packet_hash') != digest(packet.model_dump())):
        raise ValueError('Final export requires passing validation and explicit approval')
    if not draft:
        require_clean_final_text(packet, career)
    out = Path(directory)/run_id/('draft' if draft else 'final')
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    render_warnings = []
    for name, text in zip(('resume', 'cover_letter'), preview(packet, career)):
        if name == 'cover_letter' and not m.get('settings', {}).get('include_cover_letter', True):
            continue
        if draft:
            text = 'DRAFT — NOT APPROVED\n\n' + text
        invalid = XML_INVALID.findall(text)
        if invalid:
            render_warnings.append({'document': name, 'code': 'invalid_xml_characters', 'count': len(invalid),
                                    'codepoints': sorted({f'U+{ord(c):04X}' for c in invalid}),
                                    'action': 'replaced with visible U+FFFD in draft only; source preserved'})
            text = XML_INVALID.sub('\ufffd', text)
        md = out/f'{name}.md'
        md.write_text(text, encoding='utf-8')
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.75)
        normal = document.styles['Normal']
        normal.font.name, normal.font.size = 'Calibri', Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.05
        for level in range(1, 4):
            heading = document.styles[f'Heading {level}']
            heading.font.name = 'Calibri'
            heading.font.size = Pt({1: 18, 2: 13, 3: 11}[level])
            heading.paragraph_format.keep_with_next = True
        for line in text.split('\n\n'):
            if line.startswith('#'):
                document.add_heading(line.lstrip('# '), level=min(len(line)-len(line.lstrip('#')), 3))
            elif line.startswith('- '):
                document.add_paragraph(line[2:], style='List Bullet')
            else:
                document.add_paragraph(line)
        docx = out/f'{name}.docx'
        document.save(docx)
        paths.extend([md, docx])
    m['render_warnings'] = render_warnings
    for path in paths:
        m['artifacts'].append({'path': str(path), 'sha256': __import__('hashlib').sha256(path.read_bytes()).hexdigest(),
                               'draft': draft, 'packet_hash': digest(packet.model_dump()), 'created': now()})
    store.save(m)
    report = out/'audit.json'
    report.write_text(json.dumps(m, indent=2), encoding='utf-8')
    return paths + [report]

def export_pdf(store=None, run_id=None, paths=None, **kwargs):
    if store is None:
        import shutil
        return {'available': bool(shutil.which('libreoffice') or shutil.which('soffice')),
                'reason': 'Local LibreOffice is required for PDF rendering.'}
    from .layout import render_layout
    return render_layout(store, run_id, paths, **kwargs)
